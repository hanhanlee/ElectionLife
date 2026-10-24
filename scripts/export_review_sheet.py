"""把題庫（事件、選項、加減分）、命運牌、突發事件、劇情鏈與數值說明匯出成一份 Excel，給人工審閱用。

內容檔改了之後重跑一次即可：python -I scripts/export_review_sheet.py
輸出：docs/選舉人生_題庫總表.xlsx
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "web" / "src" / "content"
GAME = ROOT / "web" / "src" / "game"
OUT = ROOT / "docs" / "選舉人生_題庫總表.xlsx"

STATS = ["approval", "media", "treasury", "achievement", "council", "team", "business", "party", "integrity"]
STAT_LABELS = {
    "approval": "民調支持", "media": "媒體聲量", "treasury": "市庫財政", "achievement": "施政成果", "council": "議會關係",
    "team": "團隊", "business": "政商關係", "party": "政黨支持", "integrity": "原則（隱藏）",
}
PHASES = {
    "launch": "參選", "platform": "主打政見", "campaign": "選戰", "eve": "選前之夜",
    "term": "任期", "incumbency": "任期尾聲（現任限定）", "reelection": "連任選戰",
}
STAGE = {"launch": "第一階段", "platform": "第一階段", "campaign": "第一階段", "eve": "第一階段",
         "term": "第二階段", "incumbency": "第二階段", "reelection": "第二階段"}
FATE_KINDS = {"opportunity": "機運", "contract": "契約", "ordeal": "試煉"}
FATE_STAGES = {"campaign": "選戰", "term": "任期"}

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")


def load(name):
    return json.loads((CONTENT / name).read_text(encoding="utf-8"))


def ts_record(path, const):
    """從 TypeScript 檔讀出 `export const NAME ... = { key: value, ... };` 的簡單物件"""
    text = (GAME / path).read_text(encoding="utf-8")
    block = re.search(rf"export const {const}\b[^=]*=\s*\{{(.*?)\n\}};", text, re.S).group(1)
    out = {}
    for m in re.finditer(r'^\s*(\w+):\s*("(?:[^"\\]|\\.)*"|-?[\d.]+)', block, re.M):
        v = m.group(2)
        out[m.group(1)] = json.loads(v) if v.startswith('"') else float(v)
    return out


def cond(c):
    if not c:
        return ""
    parts = []
    if c.get("flags"):
        parts.append("需要 " + " + ".join(c["flags"]))
    if c.get("anyFlags"):
        parts.append("任一 " + " / ".join(c["anyFlags"]))
    if c.get("notFlags"):
        parts.append("排除 " + " / ".join(c["notFlags"]))
    for k, v in (c.get("min") or {}).items():
        parts.append(f"{STAT_LABELS.get(k, k)} ≥ {v}")
    for k, v in (c.get("max") or {}).items():
        parts.append(f"{STAT_LABELS.get(k, k)} ≤ {v}")
    return "；".join(parts)


def delta(d):
    if not d:
        return ""
    return "、".join(f"{STAT_LABELS.get(k, k)} {v:+g}" for k, v in d.items())


def arcs_text(arcs):
    out = []
    for a in arcs or []:
        bits = [a["id"]]
        if a.get("open"):
            bits.append("開啟")
        if a.get("status"):
            bits.append(f"→{a['status']}")
        if a.get("add"):
            bits.append("+" + ",".join(f"{k}{v:+g}" for k, v in a["add"].items()))
        if a.get("set"):
            bits.append("設定 " + ",".join(f"{k}={v}" for k, v in a["set"].items()))
        out.append(" ".join(bits) + f"（{a['note']}）")
    return "\n".join(out)


def write_sheet(ws, headers, rows, widths, review_col=None, number_cols=()):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1] if c - 1 < len(widths) else 12
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if cell.column in number_cols else None)
            if review_col and cell.column == review_col:
                cell.fill = REVIEW_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    events, fates, crises, arcs = load("events.json"), load("fates.json"), load("crises.json"), load("arcs.json")
    hints = ts_record("types.ts", "STAT_HINTS")
    weights = ts_record("model.ts", "VOTE_WEIGHTS")
    notes = ts_record("model.ts", "WEIGHT_NOTES")
    start = ts_record("model.ts", "START_STATS")
    arc_title = {a["id"]: a["title"] for a in arcs}

    wb = Workbook()

    # ---- 說明 ----
    ws = wb.active
    ws.title = "說明"
    lines = [
        ("《選舉人生：市長篇》題庫總表", ""),
        ("用途", "給人工審閱：題目情境是否合理、選項之間的取捨是否清楚、加減分方向與幅度是否說得通。"),
        ("審閱方式", "每張表最後一欄「審閱備註」（黃色）可以直接寫意見；不要改其他欄，內容以遊戲的 JSON 檔為準，本表由 scripts/export_review_sheet.py 自動產生。"),
        ("工作表", "數值說明：九個數值的意義、起始值、對得票的權重。\n事件與選項：每個選項一列，九個數值各一欄（正數加分、負數扣分）。\n命運牌：每階段第 3、6、9 回合後的三選一。\n突發事件：數值越界時自動觸發，不是選項。\n劇情鏈：會在之後回合找上門的帳。"),
        ("階段", "第一階段（競選）12 回合：參選 1、主打政見 1、選戰 9、選前之夜 1，結束時第一次投票。\n第二階段（任期與連任）12 回合：任期 8、任期尾聲 3、連任選戰 1，結束時連任投票。"),
        ("表上的數字是「原始效果」", "遊戲實際套用時還會調整：\n1. 民調邊際遞減：民調超過 50 後，民調加分打折（80 時剩七成、100 時剩一半）。\n2. 團隊倍率：施政成果的加分乘上團隊係數（團隊高做得出成果）。\n3. 城市選民結構：同一個議題在不同城市，民調效果放大或縮小。\n4. 資質：選項標了「能力」時，擅長的資質額外加分、不擅長的扣分。\n5. 現任資源的選民疲乏：每多動用一次，民調加分打 0.65 折。\n6. 當選後兌現：挑戰者在競選期的承諾，要當選後才影響施政成果與市庫。\n7. 命運牌：機運與契約的好處 ×0.4、代價 ×1.5、拿到的經費 ×0.6；試煉成功 ×0.75、失敗 ×1.35。"),
        ("每回合固定變化", "媒體聲量每回合退燒 3 點（最低 20）。競選回合依競選規模付人事、廣告、活動；任期回合付服務處開銷 200 萬。每 3 回合入帳一次。"),
        ("選舉模型", "勝選機率 = sigmoid（基準 + Σ 權重 ×（數值 − 50）/ 10 − 對手強度 + 選前動員），投票時再加一點隨機。第一次投票的基準 −0.32，連任投票 +0.5。權重見「數值說明」。"),
        ("旗標與條件", "需要 X：有這個旗標才會出現；排除 X：有這個旗標就不會出現。party-*：政黨；bg-*：出身；city-*：城市；trait-*：資質；opp-*：對手；promise-*：承諾過的政見；arc-*：劇情鏈狀態。"),
    ]
    for k, v in lines:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws["A1"].font = Font(name=FONT, bold=True, size=14)

    # ---- 數值說明 ----
    ws = wb.create_sheet("數值說明")
    rows = []
    for k in STATS:
        rows.append([STAT_LABELS[k], k, hints.get(k, ""), int(start.get(k, 0)), weights.get(k, 0), notes.get(k, ""), ""])
    write_sheet(ws, ["數值", "程式代號", "意義與作用", "起始值（政黨、出身、資質、城市會再加減）", "對得票的權重（每 10 點，logit）", "權重說明", "審閱備註"],
                rows, [14, 14, 60, 18, 18, 50, 40], review_col=7, number_cols=(4, 5))

    # ---- 事件與選項 ----
    ws = wb.create_sheet("事件與選項")
    headers = ["事件ID", "階段", "回合類型", "事件標題", "情境", "出現條件", "抽選權重", "劇情鏈", "選項ID", "選項文字", "結果文字"]
    headers += [STAT_LABELS[k] for k in STATS]
    headers += ["經費（萬）", "當選後兌現", "延遲後果", "設定旗標", "選項條件", "能力", "違法風險", "動用現任資源", "政治帳本", "來源狀態", "真實原型", "審閱備註"]
    rows = []
    for e in events:
        for i, o in enumerate(e["options"]):
            ocond = []
            if o.get("requires"):
                ocond.append("鎖住除非 " + cond(o["requires"]))
            if o.get("appearsIf"):
                ocond.append("出現條件 " + cond(o["appearsIf"]))
            d = o.get("delayed")
            el = o.get("onElected")
            rows.append([
                e["id"], STAGE[e["phase"]], PHASES[e["phase"]], e["title"], e["scenario"],
                cond(e.get("conditions")), e.get("weight", 1), arc_title.get(e.get("arc"), ""),
                o["id"], o["label"], o["outcome"],
                *[o["effects"].get(k) for k in STATS],
                o.get("funds"),
                f"{delta(el['effects'])}（{el['text']}）" if el else "",
                f"{d['afterTurns']} 回合後：{delta(d['effects'])}（{d['text']}）" if d else "",
                ", ".join(o.get("setFlags", [])), "；".join(ocond), o.get("skill", ""), o.get("legalRisk", ""),
                "是" if o.get("perk") else "", arcs_text(o.get("arcs")),
                e["source"]["status"], e["source"].get("realCase") or e["source"].get("prototype", ""), "",
            ])
    widths = [16, 10, 12, 18, 44, 24, 8, 12, 12, 30, 40] + [8] * len(STATS) + [9, 30, 34, 18, 22, 10, 10, 8, 30, 10, 36, 40]
    first_stat = 12
    write_sheet(ws, headers, rows, widths, review_col=len(headers),
                number_cols=tuple(range(first_stat, first_stat + len(STATS) + 1)) + (7,))

    # ---- 命運牌 ----
    ws = wb.create_sheet("命運牌")
    headers = ["牌ID", "類型", "階段", "牌名", "情境", "牌面提示", "出現條件", "結果"]
    headers += [STAT_LABELS[k] for k in STATS]
    headers += ["經費（萬）", "加成", "持續效果", "護盾", "延遲後果", "政治帳本", "試煉成功率（不顯示給玩家）", "結果文字", "審閱備註"]
    rows = []

    def fate_row(c, label, eff, chance=""):
        b, p, s, d = eff.get("boost"), eff.get("passive"), eff.get("shield"), eff.get("delayed")
        texts = eff["text"] if isinstance(eff["text"], list) else [eff["text"]]
        return [
            c["id"], FATE_KINDS[c["kind"]], " / ".join(FATE_STAGES[x] for x in c["stages"]), c["title"], c["scene"], c["hint"],
            cond(c.get("conditions")), label,
            *[eff.get("effects", {}).get(k) for k in STATS],
            eff.get("funds"),
            f"{'/'.join(b.get('skills', []) + b.get('issues', []))}：{b['charges']} 次、{b['turns']} 回合內（{b['label']}）" if b else "",
            f"每回合 {delta(p['effects'])}，{p['turns']} 回合（{p['label']}）" if p else "",
            f"{s['turns']} 回合內擋下一次危機的一半衝擊（{s['label']}）" if s else "",
            f"{d['afterTurns']} 回合後：{delta(d['effects'])}（{d['text']}）" if d else "",
            arcs_text(eff.get("arcs")), chance, "\n".join(texts), "",
        ]

    for c in fates:
        if c.get("ordeal"):
            o = c["ordeal"]
            mods = "；".join(f"{cond(m['when'])} {m['add']:+.2f}" for m in o["mods"])
            chance = f"基礎 {o['base']:.2f}" + (f"；{mods}" if mods else "") + "（限 0.20–0.80）"
            rows.append(fate_row(c, "成功", o["success"], chance))
            rows.append(fate_row(c, "失敗", o["failure"]))
        else:
            rows.append(fate_row(c, "—", c["effect"]))
    widths = [16, 8, 10, 18, 40, 40, 20, 8] + [8] * len(STATS) + [9, 30, 26, 24, 34, 30, 30, 44, 40]
    write_sheet(ws, headers, rows, widths, review_col=len(headers), number_cols=tuple(range(9, 9 + len(STATS) + 1)))

    # ---- 突發事件 ----
    ws = wb.create_sheet("突發事件")
    headers = ["ID", "名稱", "好壞", "觸發條件", "階段", "內容"] + [STAT_LABELS[k] for k in STATS] + ["經費（萬）", "餘波", "設定旗標", "可重複", "審閱備註"]
    rows = []
    for c in sorted(crises, key=lambda c: c["priority"]):
        a = c.get("aftermath")
        rows.append([
            c["id"], c["title"], "負面" if c["tone"] == "bad" else "正面", cond(c["when"]),
            " / ".join(PHASES[p] for p in c.get("phases", [])) or "全程", c["text"],
            *[c["effects"].get(k) for k in STATS], c.get("funds"),
            f"{a['turns']} 回合，每回合 {delta(a['effects'])}" if a else "", ", ".join(c.get("setFlags", [])),
            "是" if c.get("repeat") else "", "",
        ])
    write_sheet(ws, headers, rows, [16, 16, 6, 24, 20, 50] + [8] * len(STATS) + [9, 26, 16, 6, 40],
                review_col=len(headers), number_cols=tuple(range(7, 7 + len(STATS) + 1)))

    # ---- 劇情鏈 ----
    ws = wb.create_sheet("劇情鏈")
    headers = ["ID", "名稱", "種類", "帳本進度", "狀態說明", "節點（依序）", "審閱備註"]
    rows = []
    title_of = {e["id"]: e["title"] for e in events}
    for a in arcs:
        nodes = "\n".join(
            f"{i + 1}. {title_of.get(n['event'], n['event'])}（{PHASES[n['phase']]}"
            + (f"，只在 {'/'.join(n['when'])} 時" if n.get("when") else "") + "）"
            for i, n in enumerate(a["nodes"]))
        rows.append([a["id"], a["title"], {"promise": "政見承諾", "debt": "人情債", "scandal": "爭議"}[a["kind"]], a["progress"],
                     "\n".join(f"{k}：{v}" for k, v in a["statusLabels"].items()), nodes, ""])
    write_sheet(ws, headers, rows, [14, 14, 10, 40, 36, 50, 40], review_col=len(headers))

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    n_opts = sum(len(e["options"]) for e in events)
    print(f"寫入 {OUT.relative_to(ROOT)}：{len(events)} 事件、{n_opts} 選項、{len(fates)} 張命運牌、{len(crises)} 突發事件、{len(arcs)} 條劇情鏈")
    return 0


if __name__ == "__main__":
    sys.exit(main())
